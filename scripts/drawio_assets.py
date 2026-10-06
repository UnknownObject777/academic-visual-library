"""Small renderer/exporter for the bundled primitive-based template library.

Reads final draw.io XML directly, not an intermediate scene schema. Unsupported
shapes are reported instead of silently converted. Original draw.io is canonical.
"""
from pathlib import Path
import argparse, html, math, re, xml.etree.ElementTree as ET

def load(path):
    root = ET.parse(path).getroot()
    if root.tag == 'mxfile':
        root = root.find('diagram/mxGraphModel')
    if root is None:
        raise ValueError('Uncompressed mxGraphModel required; export uncompressed XML.')
    cells = {c.get('id'): c for c in root.iter('mxCell')}
    def position(c, seen=()):
        if c.get('id') in seen:
            raise ValueError('Cyclic cell parenting')
        g = c.find('mxGeometry')
        x, y = (float(g.get('x',0)), float(g.get('y',0))) if g is not None else (0,0)
        parent = cells.get(c.get('parent'))
        if parent is not None and parent.get('vertex') == '1':
            px,py,_,_ = position(parent,seen+(c.get('id'),))
            x,y = x+px,y+py
        return x,y,float(g.get('width',0)) if g is not None else 0,float(g.get('height',0)) if g is not None else 0
    nodes, edges = [], []
    for c in cells.values():
        st = dict(p.split('=',1) for p in c.get('style','').split(';') if '=' in p)
        if c.get('style','').startswith('group;'):
            continue
        text = re.sub(r'<br\s*/?>','\n',c.get('value',''),flags=re.I)
        text = html.unescape(re.sub('<[^>]*>','',text))
        if c.get('vertex') == '1':
            shape = st.get('shape','ellipse' if 'ellipse;' in c.get('style','') else 'rect')
            if shape not in ('rect','cube','ellipse','trapezoid','rhombus'):
                raise ValueError('Unsupported shape: '+shape)
            nodes.append((c.get('id'),position(c),st,text,shape))
        elif c.get('edge') == '1':
            edges.append((c,st,text))
    byid = {n[0]:n for n in nodes}
    paths = []
    for c,st,text in edges:
        g=c.find('mxGeometry')
        parent=cells.get(c.get('parent'))
        ox,oy,_,_=position(parent) if parent is not None else (0,0,0,0)
        points=[]
        for end,attr in [('source','exit'),('target','entry')]:
            if c.get(end) in byid:
                _,(x,y,w,h),_,_,_=byid[c.get(end)]
                point=(x+w*float(st.get(attr+'X',.5)),y+h*float(st.get(attr+'Y',.5)))
            else:
                p=g.find("mxPoint[@as='"+end+"Point']") if g is not None else None
                if p is None:
                    raise ValueError('Unresolved edge endpoint '+str(c.get('id')))
                point=(float(p.get('x',0))+ox,float(p.get('y',0))+oy)
            points.append(point)
        mid=[]
        if g is not None:
            mid=[(float(p.get('x',0))+ox,float(p.get('y',0))+oy) for p in g.findall("Array[@as='points']/mxPoint")]
        pts=[points[0]]+mid+[points[1]]
        paths.append((pts,st,text))
    coords=[(x,y) for _,(a,b,w,h),_,_,_ in nodes for x,y in [(a,b),(a+w,b+h)]]
    coords += [p for pts,_,_ in paths for p in pts]
    xmin=min(x for x,y in coords)-20; ymin=min(y for x,y in coords)-20
    xmax=max(x for x,y in coords)+20; ymax=max(y for x,y in coords)+20
    return nodes,paths,(xmin,ymin,xmax-xmin,ymax-ymin)

def svg(path, out):
    nodes,paths,bounds=load(path); xmin,ymin,w,h=bounds
    chunks=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{xmin} {ymin} {w} {h}" width="{w}" height="{h}">',f'<rect x="{xmin}" y="{ymin}" width="{w}" height="{h}" fill="white"/>']
    for i,(pts,st,text) in enumerate(paths):
        color=st.get('strokeColor','#475569'); arrow=st.get('endArrow','classic') not in ('none','')
        chunks.append(f'<defs><marker id="a{i}" markerWidth="8" markerHeight="6" refX="7" refY="3" orient="auto"><path d="M0,0 L8,3 L0,6 Z" fill="{color}"/></marker></defs>')
        d='M '+' L '.join(f'{x},{y}' for x,y in pts)
        dash=' stroke-dasharray="5 4"' if st.get('dashed')=='1' else ''
        marker=f' marker-end="url(#a{i})"' if arrow else ''
        chunks.append(f'<path d="{d}" fill="none" stroke="{color}" stroke-width="{st.get("strokeWidth",1.5)}"{dash}{marker}/>')
        if text:
            x,y=pts[len(pts)//2]; chunks.append(f'<text x="{x}" y="{y-4}" font-family="Arial" font-size="11">{html.escape(text)}</text>')
    edge_end=len(chunks); underlays=[]
    for ident,(x,y,w,h),st,text,shape in nodes:
        start=len(chunks)
        fill=st.get('fillColor','#FFFFFF'); stroke=st.get('strokeColor','#64748B')
        common=f'fill="{fill}" stroke="{stroke}" stroke-width="{st.get("strokeWidth",1.5)}"'
        if shape=='ellipse':
            chunks.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" {common}/>')
        elif shape in ('trapezoid','rhombus'):
            pts=[(x+w*.2,y),(x+w*.8,y),(x+w,y+h),(x,y+h)] if shape=='trapezoid' else [(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)]
            chunks.append('<polygon points="'+' '.join(f'{a},{b}' for a,b in pts)+f'" {common}/>')
        else:
            radius=6 if st.get('rounded')=='1' else 0
            chunks.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{radius}" {common}/>')
            if shape=='cube':
                d=min(w,h)*.15
                chunks.append(f'<path d="M{x},{y+d} H{x+w-d} V{y+h} M{x+w-d},{y+d} L{x+w},{y} M{x+w-d},{y+h} L{x+w},{y+h-d}" fill="none" stroke="{stroke}"/>')
        lines=text.splitlines(); size=float(st.get('fontSize',12)); y0=y+h/2-(len(lines)-1)*size*.6
        if st.get('verticalAlign')=='top': y0=y+size*.8
        elif st.get('verticalAlign')=='bottom': y0=y+h-len(lines)*size*1.2+size*.6
        for j,line in enumerate(lines):
            chunks.append(f'<text x="{x+w/2}" y="{y0+j*size*1.2}" dominant-baseline="middle" text-anchor="middle" font-family="Arial, sans-serif" font-size="{size}" fill="{st.get("fontColor","#24323D")}">{html.escape(line)}</text>')
        contains=sum(other!=ident and a>=x and b>=y and a+ow<=x+w and b+oh<=y+h for other,(a,b,ow,oh),_,_,_ in nodes)
        if contains>=2:
            underlays.extend(chunks[start:]); del chunks[start:]
    chunks=chunks[:2]+underlays+chunks[2:edge_end]+chunks[edge_end:]
    chunks.append('</svg>'); Path(out).write_text('\n'.join(chunks),encoding='utf-8')

def add_ppt_slide(prs, path):
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    nodes,paths,(xmin,ymin,w,h)=load(path); slide=prs.slides.add_slide(prs.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(.25),Inches(.05),Inches(12.5),Inches(.35)); title.text=Path(path).stem
    scale=min(12.6/w,6.4/h); dx=(13.333-w*scale)/2-xmin*scale; dy=.65+(6.4-h*scale)/2-ymin*scale
    def color(value,fallback):
        return RGBColor.from_string(value.lstrip('#')) if re.fullmatch('#?[0-9A-Fa-f]{6}',value) else RGBColor.from_string(fallback)
    types={'rect':MSO_AUTO_SHAPE_TYPE.RECTANGLE,'cube':MSO_AUTO_SHAPE_TYPE.CUBE,'ellipse':MSO_AUTO_SHAPE_TYPE.OVAL,'trapezoid':MSO_AUTO_SHAPE_TYPE.TRAPEZOID,'rhombus':MSO_AUTO_SHAPE_TYPE.DIAMOND}
    connections=[]
    for pts,st,text in paths:
        for a,b in zip(pts,pts[1:]):
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]*scale+dx),Inches(a[1]*scale+dy),Inches(b[0]*scale+dx),Inches(b[1]*scale+dy))
            line.line.color.rgb=color(st.get('strokeColor',''),'475569'); line.line.width=Pt(1); connections.append(line)
        if st.get('endArrow','classic') not in ('none',''):
            tail=OxmlElement('a:tailEnd'); tail.set('type','triangle'); line._element.spPr.get_or_add_ln().append(tail)
        if text:
            a=pts[len(pts)//2]; label=slide.shapes.add_textbox(Inches(a[0]*scale+dx),Inches(a[1]*scale+dy-.2),Inches(1.5),Inches(.3)); label.text=text
    for ident,(x,y,nw,nh),st,text,shape in nodes:
        kind=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if shape=='rect' and st.get('rounded')=='1' else types[shape]
        obj=slide.shapes.add_shape(kind,Inches(x*scale+dx),Inches(y*scale+dy),Inches(nw*scale),Inches(nh*scale))
        for effect in obj._element.xpath('.//a:effectRef'): effect.set('idx','0')
        obj._element.spPr.append(OxmlElement('a:effectLst'))
        if st.get('fillColor')=='none': obj.fill.background()
        else: obj.fill.solid(); obj.fill.fore_color.rgb=color(st.get('fillColor',''),'FFFFFF')
        obj.line.color.rgb=color(st.get('strokeColor',''),'64748B'); obj.text=text
        obj.text_frame.vertical_anchor={'top':MSO_ANCHOR.TOP,'bottom':MSO_ANCHOR.BOTTOM}.get(st.get('verticalAlign'),MSO_ANCHOR.MIDDLE)
        obj.text_frame.margin_left=obj.text_frame.margin_right=Inches(.02)
        for p in obj.text_frame.paragraphs:
            p.alignment=PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name='Arial'; run.font.size=Pt(max(5,min(16,float(st.get('fontSize',12))*scale*72)))
                run.font.color.rgb=color(st.get('fontColor',''),'24323D')
        contains=sum(other!=ident and a>=x and b>=y and a+ow<=x+nw and b+oh<=y+nh for other,(a,b,ow,oh),_,_,_ in nodes)
        if contains>=2:
            tree=slide.shapes._spTree; tree.remove(obj._element); tree.insert(2,obj._element)
    return slide

def ppt(paths,out):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    for path in paths: add_ppt_slide(prs,path)
    prs.save(out)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',type=Path); parser.add_argument('--svg',type=Path); parser.add_argument('--pptx',type=Path)
    args=parser.parse_args()
    if args.svg: svg(args.input,args.svg)
    if args.pptx: ppt([args.input],args.pptx)
    if not args.svg and not args.pptx: parser.error('Select --svg and/or --pptx')
