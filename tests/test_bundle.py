from pathlib import Path
import hashlib, importlib.util, json, re, tempfile, unittest
import xml.etree.ElementTree as ET
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('drawio_assets',ROOT/'scripts/drawio_assets.py')
renderer=importlib.util.module_from_spec(spec); spec.loader.exec_module(renderer)

class BundleTests(unittest.TestCase):
    def test_templates_and_edges_resolve(self):
        models=list((ROOT/'assets/drawio/models').glob('*.drawio'))
        self.assertEqual(len(models),37)
        for model in models:
            nodes,edges,bounds=renderer.load(model)
            self.assertTrue(nodes,model.name)
            self.assertGreater(bounds[2],0); self.assertGreater(bounds[3],0)
            self.assertTrue(all(len(points)>=2 for points,_,_ in edges))

    def test_ppt_contains_native_objects(self):
        for name,count in [('modules.pptx',1),('network-models.pptx',37)]:
            deck=Presentation(ROOT/'assets/pptx'/name)
            self.assertEqual(len(deck.slides),count)
            for slide in deck.slides:
                self.assertGreater(len(slide.shapes),1)
                self.assertFalse(any(s.shape_type==13 for s in slide.shapes))

    def test_renderer_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            source=ROOT/'assets/drawio/modules.drawio'
            renderer.svg(source,Path(directory)/'modules.svg')
            renderer.ppt([source],Path(directory)/'modules.pptx')
            svg=(Path(directory)/'modules.svg').read_text(encoding='utf-8')
            self.assertIn('Convolution',svg); self.assertIn('Attention',svg)
            texts=[s.text for s in Presentation(Path(directory)/'modules.pptx').slides[0].shapes if s.has_text_frame]
            self.assertIn('Convolution',texts); self.assertIn('Attention',texts)

    def test_unknown_shapes_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'unknown.drawio'
            graph=ET.Element('mxGraphModel'); root=ET.SubElement(graph,'root')
            cell=ET.SubElement(root,'mxCell',id='x',vertex='1',style='shape=unsupported;')
            ET.SubElement(cell,'mxGeometry',width='20',height='20')
            ET.ElementTree(graph).write(path,encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'Unsupported shape'): renderer.load(path)

    def test_provenance(self):
        for record in json.loads((ROOT/'assets/provenance.json').read_text(encoding='utf-8')):
            asset=ROOT/record.get('asset','assets/drawio/nn-model-library.xml')
            self.assertEqual(hashlib.sha256(asset.read_bytes()).hexdigest(),record['sha256'])
            self.assertTrue((ROOT/record['license_file']).is_file())

    def test_document_links(self):
        for file in ROOT.rglob('*.md'):
            for link in re.findall(r'\]\(([^)]+)\)',file.read_text(encoding='utf-8')):
                if link.startswith(('http:','https:','#')): continue
                self.assertTrue((file.parent/link.split('#')[0]).exists(),str(file)+': '+link)

if __name__=='__main__': unittest.main()
